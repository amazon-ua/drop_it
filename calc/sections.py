"""База сечений профильных труб из прайса (input/price_profile_tubes.xlsx).

Геометрические характеристики считаются по средней линии стенки
с закруглёнными углами (наружный радиус 2t, внутренний t — как у
гнутых труб по ДСТУ EN 10219 / ГОСТ 30245).
"""
from __future__ import annotations

import math
import re
from dataclasses import dataclass, replace
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
PRICE_XLSX = ROOT / "input" / "price_profile_tubes.xlsx"
# полный прайс с сайта поставщика (28.09): та же номенклатура + длина хлыста L по каждой позиции
PRICE_FULL_XLSX = ROOT / "input" / "price_profile_tubes_full.xlsx"

RHO = 7850.0  # кг/м3


@dataclass(frozen=True)
class Section:
    name: str          # обозначение, напр. "80×40×2"
    h: float           # высота, м (в плоскости сильной оси)
    b: float           # ширина, м
    t: float           # фактическая толщина стенки, м
    price: float       # грн/м
    source: str        # строка прайса
    A: float = 0.0     # м2
    Iy: float = 0.0    # м4, изгиб в плоскости h (сильная ось)
    Iz: float = 0.0    # м4, изгиб в плоскости b
    J: float = 0.0     # м4, кручение (Бредт)
    Wy: float = 0.0    # м3
    Wz: float = 0.0    # м3
    iy: float = 0.0    # м
    iz: float = 0.0    # м
    mass: float = 0.0  # кг/м
    perim: float = 0.0  # наружный периметр, м (для окраски и сварки)

    @property
    def c_t_max(self) -> float:
        """Гибкость наибольшей плоской стенки c/t (c = b − 3t, как в EN 1993-1-1)."""
        return (max(self.h, self.b) - 3 * self.t) / self.t

    @property
    def c_t_h(self) -> float:
        return (self.h - 3 * self.t) / self.t

    @property
    def c_t_b(self) -> float:
        return (self.b - 3 * self.t) / self.t


def _props(h: float, b: float, t: float) -> dict:
    """Тонкостенные характеристики по средней линии с закруглёнными углами."""
    rm = 1.5 * t              # радиус средней линии угла (ro=2t, ri=t)
    hm, bm = h - t, b - t     # размеры по средней линии
    # прямые участки
    lh = hm - 2 * rm          # вертикальные стенки (вдоль h)
    lb = bm - 2 * rm          # горизонтальные стенки (вдоль b)
    arc = 0.5 * math.pi * rm  # четверть окружности
    p_mid = 2 * lh + 2 * lb + 4 * arc
    A = p_mid * t
    # моменты инерции относительно оси, параллельной b (изгиб в плоскости h)
    # вертикальные стенки
    Iy = 2 * (t * lh ** 3 / 12)
    # горизонтальные стенки на расстоянии hm/2
    Iy += 2 * (lb * t) * (hm / 2) ** 2 + 2 * (lb * t ** 3 / 12)
    # углы: четверть кольца, центр на (lb/2, lh/2)
    # момент инерции дуги радиуса r вокруг собственной оси центра: для четверти
    # ∫ (yc + r sinθ)^2 r t dθ, θ∈[0, π/2]
    def arc_I(offset: float) -> float:
        r = rm
        # ∫0^{π/2} (offset + r sinθ)^2 dθ * r t
        integ = (offset ** 2) * (math.pi / 2) + 2 * offset * r * 1.0 + r ** 2 * (math.pi / 4)
        return integ * r * t
    Iy += 4 * arc_I(lh / 2)
    Iz = 2 * (t * lb ** 3 / 12)
    Iz += 2 * (lh * t) * (bm / 2) ** 2 + 2 * (lh * t ** 3 / 12)
    Iz += 4 * arc_I(lb / 2)
    # кручение: Бредт
    Am = hm * bm - (4 - math.pi) * rm ** 2
    J = 4 * Am ** 2 * t / p_mid
    Wy = Iy / (h / 2)
    Wz = Iz / (b / 2)
    perim_out = 2 * (h + b) - (8 - 2 * math.pi) * (2 * t)
    return dict(A=A, Iy=Iy, Iz=Iz, J=J, Wy=Wy, Wz=Wz,
                iy=math.sqrt(Iy / A), iz=math.sqrt(Iz / A),
                mass=A * RHO, perim=perim_out)


_NUM = r"(\d+(?:[.,]\d+)?)"


def _parse_row(name: str):
    """'Труба профільна 80*40*3,0(2,8) г/к' -> (h, b, t_actual) в мм."""
    s = name.replace(" ", "")
    m = re.search(r"(\d+)\*(\d+)\*\(?(\d+(?:,\d+)?)\)?(?:\((\d+(?:,\d+)?)\))?", s)
    if not m:
        return None
    h, b = float(m.group(1)), float(m.group(2))
    t_nom = float(m.group(3).replace(",", "."))
    t_alt = m.group(4)
    # формат "(3,0)2,8" — фактическая толщина вне скобок
    m2 = re.search(r"\*\((\d+(?:,\d+)?)\)(\d+(?:,\d+)?)", s)
    if m2:
        t_act = float(m2.group(2).replace(",", "."))
    elif t_alt:
        # "2,0(1,8)" -> фактическая 1,8; "2,0(2,2)" -> берём меньшую (консервативно)
        t_act = min(t_nom, float(t_alt.replace(",", ".")))
    else:
        t_act = t_nom
    return h, b, t_act


def load_price_list(min_t: float = 2.0, allow_substandard: bool = False,
                    min_size: float = 20.0):
    """Сечения из прайса: прямоугольные/квадратные, t_факт >= min_t мм.

    Некондиция по умолчанию исключена. При нескольких позициях с одинаковым
    сечением берётся самая дешёвая.
    """
    wb = openpyxl.load_workbook(PRICE_XLSX, data_only=True)
    ws = wb.active
    best: dict[tuple, Section] = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or row[1] is None or row[2] is None:
            continue
        name = str(row[1])
        low = name.lower()
        if "трикут" in low or "овальн" in low:
            continue
        if not allow_substandard and "некон" in low:
            continue
        parsed = _parse_row(name)
        if not parsed:
            continue
        h, b, t = parsed
        if t + 1e-9 < min_t or min(h, b) < min_size:
            continue
        price = float(row[2])
        key = (h, b, t)
        if key in best and best[key].price <= price:
            continue
        label = f"{h:g}×{b:g}×{t:g}".replace(".", ",")
        p = _props(h / 1000, b / 1000, t / 1000)
        best[key] = Section(label, h / 1000, b / 1000, t / 1000, price, name, **p)
    secs = sorted(best.values(), key=lambda s: s.price)
    return secs


def by_name(secs, name: str) -> Section:
    for s in secs:
        if s.name == name:
            return s
    raise KeyError(name)


if __name__ == "__main__":
    secs = load_price_list()
    print(f"{len(secs)} сечений (t>=2 мм, без некондиции)")
    for s in secs:
        print(f"{s.name:>12} {s.price:7.1f} грн/м  {s.mass:5.2f} кг/м  A={s.A*1e4:5.2f} см2 "
              f"Wy={s.Wy*1e6:6.2f} Wz={s.Wz*1e6:6.2f} см3 iy={s.iy*100:4.2f} iz={s.iz*100:4.2f} см "
              f"грн/кг={s.price/s.mass:5.1f}  c/t={s.c_t_max:4.1f}")


def _norm_name(name: str) -> str:
    s = re.sub(r";\s*L\s*-.*$", "", str(name))
    s = s.replace(" (кількість обмежена)", "")
    return re.sub(r"\s+", " ", s).strip()


_BARS = None


def bar_lengths_from_price(source: str):
    """Длины хлыстов позиции по полному прайсу (колонка «Длина хлыста L, м»), м.
    Обрезки «1-3» и короткие остатки (< 5 м) не учитываются. None — позиции нет в полном прайсе."""
    global _BARS
    if _BARS is None:
        _BARS = {}
        if PRICE_FULL_XLSX.exists():
            ws = openpyxl.load_workbook(PRICE_FULL_XLSX, data_only=True).active
            for r in list(ws.iter_rows(values_only=True))[1:]:
                if not r[1] or r[2] is None:
                    continue
                txt = re.sub(r"\d+\s*-\s*\d+", "", str(r[2]))          # «1-3» — обрезки
                vals = {float(v.replace(",", ".")) for v in re.findall(r"\d+(?:,\d+)?", txt)}
                vals = tuple(sorted(v for v in vals if v >= 5.0))
                if vals:
                    _BARS[_norm_name(r[1])] = vals
    return _BARS.get(_norm_name(source))


def plate_section(b: float, t: float, name: str | None = None) -> Section:
    """Полоса b × t (b — поперёк стержня, плашмя): характеристики как у сплошного прямоугольника."""
    A = b * t
    Iy, Iz = b * t ** 3 / 12, t * b ** 3 / 12
    return Section(name or f"полоса {b * 1000:g}×{t * 1000:g}", t, b, t, 0.0, "полоса", A=A, Iy=Iy, Iz=Iz,
                   J=b * t ** 3 / 3, Wy=Iy / (t / 2), Wz=Iz / (b / 2), iy=math.sqrt(Iy / A),
                   iz=math.sqrt(Iz / A), mass=A * RHO, perim=2 * (b + t))


def reinforced(base: Section, add: Section, name: str | None = None) -> Section:
    """Труба base с приваренным вплотную снизу усилением add (труба или полоса, add.h — по высоте base).
    Центр тяжести смещается вниз; Wy — по наиболее удалённой фибре (верх base)."""
    y2 = -(base.h / 2 + add.h / 2)
    A = base.A + add.A
    yc = add.A * y2 / A
    Iy = base.Iy + base.A * yc ** 2 + add.Iy + add.A * (y2 - yc) ** 2
    Iz = base.Iz + add.Iz
    cmax = max(base.h / 2 - yc, base.h / 2 + add.h + yc)
    return replace(base, name=name or f"{base.name}+{add.name}", A=A, Iy=Iy, Iz=Iz, J=base.J + add.J,
                   Wy=Iy / cmax, Wz=Iz / (max(base.b, add.b) / 2), iy=math.sqrt(Iy / A),
                   iz=math.sqrt(Iz / A), mass=base.mass + add.mass)


def with_gussets(base: Section, tg: float = 0.005, hg: float | None = None, name: str | None = None) -> Section:
    """Стропило в узле А между двумя фасонками tg. Фасонки — от оси стропила вниз на hg (по нормали к оси);
    по умолчанию в запас — только нижняя половина стенки (hg = h/2)."""
    hg = base.h / 2 if hg is None else hg
    Ag = 2 * tg * hg
    yg = -hg / 2
    A = base.A + Ag
    yc = Ag * yg / A
    Iy = base.Iy + base.A * yc ** 2 + 2 * tg * hg ** 3 / 12 + Ag * (yg - yc) ** 2
    Iz = base.Iz + 2 * hg * tg * ((base.b + tg) / 2) ** 2
    return replace(base, name=name or f"{base.name}+фасонки", A=A, Iy=Iy, Iz=Iz,
                   Wy=Iy / max(base.h / 2 - yc, hg + yc), Wz=Iz / (base.b / 2 + tg), iy=math.sqrt(Iy / A),
                   iz=math.sqrt(Iz / A))
