"""Формирование результатов: output/report.md, output/smeta_optimized.xlsx, output/navis_optimized.html."""
from __future__ import annotations

import json
import math
from dataclasses import replace
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import checks as C
import costs as K
import model as M
from baseline import baseline_scheme
from drawing import build_html
from report_data import cutting_plan, member_table, piece_list, scheme_from_result

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "output"


def fmt(v, d=0):
    return f"{v:,.{d}f}".replace(",", " ")


def load_results(mode):
    p = OUT / f"opt_results_{mode}.json"
    if not p.exists():
        return []
    return [r for r in json.loads(p.read_text()) if "total" in r]


def analyse(sc, loads=None):
    an = C.Analysis(sc, loads or M.Loads()).run_all()
    return an


def main():
    loads = M.Loads()
    # --- исходный вариант
    base = baseline_scheme()
    an_b = analyse(base)
    q_b = K.quantities(an_b.model, base)
    cal = K.Calibration(q_b)
    est_b = K.estimate(q_b, cal)
    mt_b = member_table(an_b)
    # --- оптимизация
    res_all = load_results("all")
    res_norm = load_results("normal")
    best = res_all[0]
    sc = scheme_from_result(best)
    an = analyse(sc)
    pieces, q = piece_list(an.model, sc)
    est = K.estimate(q, cal)
    mt = member_table(an)
    plan = cutting_plan(pieces)
    # --- чувствительность
    sens = {}
    for tag, ld in (("Ch=0.70 (тип местности II)", replace(loads, Ch=0.70)),
                    ("Ch=0.90 (тип I, как в черновике)", replace(loads, Ch=0.90)),
                    ("снег T=100 лет (γfm=1.14, как в черновике)", replace(loads, gfm_snow=1.14))):
        a2 = C.Analysis(sc, ld).run_all()
        gu = a2.group_util()
        g, (u, gov) = max(gu.items(), key=lambda kv: kv[1][0])
        sens[tag] = (u, g, gov)
    best_norm = res_norm[0] if res_norm else None

    notes = [
        "Нагрузки: ДБН В.1.2-2:2006 (зі змінами № 1, 2). Харків: S0 = 1600 Па, W0 = 430 Па; T = 50 років "
        "(γfm = 1.00 для снігу і вітру); експлуатаційні значення при η = 0.02 (γfe = 0.49 / 0.21).",
        "Ветер: тип местности III (Ch = 0.40 при z ≤ 5 м), навес — схема 11 дод. И, тип I, α = 12°: "
        "Ce = +0.62 / −1.04 / −0.88 / −0.08 (интерполяция 10°…20°), трение Cf = 0.04.",
        "Снег: μ = 1 (схема 1 дод. Ж, α = 12°), дополнительно — снег на одном скате (п. 8.7).",
        "Сталь: Ry = 230 МПа (С235 / S235JR), γc = 1.0, γn = 1.0 (СС1). Проверки по методике ДБН В.2.6-198; "
        "расчётные длины сжатых стержней — из расчёта на устойчивость всей пространственной модели.",
        "Колонны бетонируются в лунки 300×300 глубиной 1.20–1.35 м (без изменений); заделка колонны в бетон "
        "уменьшена до 0.70 м (низ колонны заглушён, хомуты Ø8 сохранены).",
        "Цены — прайс «Профильные трубы» (по метражу, с резкой). Трубы с фактической стенкой < 2 мм и "
        "некондиция исключены.",
    ]
    title = best["label"]
    html_text = build_html(an, sc, est, est_b, pieces, title, notes)
    (OUT / "navis_optimized.html").write_text(html_text, encoding="utf-8")

    # ---------------- XLSX ----------------
    wb = Workbook()
    bold = Font(bold=True)
    head_fill = PatternFill("solid", fgColor="DDEBF7")

    def sheet(title, header, rows, widths=None):
        ws = wb.create_sheet(title)
        ws.append(header)
        for c in ws[1]:
            c.font = bold
            c.fill = head_fill
            c.alignment = Alignment(wrap_text=True, vertical="top")
        for r in rows:
            ws.append(r)
        for i, w in enumerate(widths or [], 1):
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.freeze_panes = "A2"
        return ws

    wb.remove(wb.active)
    base_lines = {n: v for n, v, k in est_b["lines"]}
    rows = []
    for n, v, k in est["lines"]:
        rows.append([n, round(base_lines.get(n, 0.0)), round(v), None,
                     {"metal": "по ведомости", "const": "без изменений", "roof": "вне оптимизации (без изменений)",
                      "area": "∝ площади окраски", "weld": "∝ длине швов", "labor": "∝ трудоёмкости"}[k]])
    ws = sheet("Сравнение", ["Статья", "Было (черновик), грн", "Стало, грн", "Δ, грн", "Как пересчитано"], rows,
               [46, 18, 14, 12, 34])
    n = len(rows) + 1
    for i in range(2, n + 1):
        ws.cell(i, 4).value = f"=C{i}-B{i}"
    ws.cell(n + 1, 1).value = "ИТОГО"
    ws.cell(n + 1, 1).font = bold
    for col in (2, 3, 4):
        L = get_column_letter(col)
        ws.cell(n + 1, col).value = f"=SUM({L}2:{L}{n})"
        ws.cell(n + 1, col).font = bold
    ws.cell(n + 3, 1).value = "Смета xlsx (исходная, для справки): итог 239 730 грн, трубы 80 460 грн."
    ws.cell(n + 4, 1).value = (f"Масса металла: было {q_b['mass']:.0f} кг, стало {q['mass']:.0f} кг; "
                               f"площадь окраски {q_b['area']:.1f} → {q['area']:.1f} м²; "
                               f"швы {q_b['weld_len']:.1f} → {q['weld_len']:.1f} м; деталей {q_b['n_cuts']} → {q['n_cuts']}.")

    rows = []
    for name, d in sorted(est["by_sec"].items(), key=lambda kv: -kv[1]["cost"]):
        rows.append([f"□{name}", ", ".join(sorted(M.GROUP_INFO[g][0] for g in d["groups"])),
                     round(d["L"], 2), d["n"], round(d["kg"], 1), d["sec"].price, round(d["cost"])])
    ws = sheet("Ведомость металла", ["Профиль", "Назначение", "Длина, м (с пропилами)", "Заготовок",
                                      "Масса, кг", "Цена, грн/м", "Стоимость, грн"], rows, [16, 50, 14, 10, 10, 11, 14])
    n = len(rows) + 1
    ws.cell(n + 1, 1).value = "ИТОГО трубы"
    for col in (3, 5, 7):
        L = get_column_letter(col)
        ws.cell(n + 1, col).value = f"=SUM({L}2:{L}{n})"
    ws.cell(n + 2, 1).value = f"Пластины (фасонки, крышки колонн): {q['plate_kg']:.1f} кг × {K.PLATE_PRICE:.0f} грн/кг"

    rows = [[p["name"], f"□{p['sec']}", p["L"], p["n"], round(p["mass"], 1),
             ("стык на раме: " + str(p["parts"]) + " заготовки") if p["parts"] > 1 else ""] for p in pieces]
    sheet("Детали", ["Деталь", "Сечение", "Длина, м", "Кол-во", "Масса, кг", "Примечание"], rows,
          [36, 16, 10, 8, 10, 30])

    rows = []
    for secname, bars in plan.items():
        for i, b in enumerate(bars, 1):
            rows.append([f"□{secname}", i, " + ".join(f"{L:.3f}" for L, _ in b["cuts"]),
                         "; ".join(sorted(set(nm for _, nm in b["cuts"]))), round(b["free"], 3)])
    sheet("Раскрой (хлыст 6 м)", ["Профиль", "Хлыст №", "Резы, м", "Детали", "Остаток, м"], rows,
          [16, 8, 40, 50, 10])

    rows = [[r["name"], f"□{r['sec']}", round(r["Nc"], 2), round(r["Nt"], 2), round(r["M"], 3),
             round(r["u"], 3), r["gov"]] for r in mt]
    sheet("Проверки", ["Элемент", "Сечение", "N сж, кН", "N раст, кН", "M, кН·м", "Коэф. исп.", "Определяющая проверка"],
          rows, [34, 16, 10, 10, 10, 10, 70])

    rows = []
    for r in res_all:
        rows.append([r["label"], round(r["total"]), round(r["pipes"]), round(r["mass"]),
                     "; ".join(f"{M.GROUP_INFO[g][0]}: □{s}" for g, s in r["groups"].items())])
    sheet("Варианты схем", ["Схема", "Итог, грн", "Трубы, грн", "Масса, кг", "Сечения"], rows, [70, 12, 12, 10, 120])

    rows = [["S0 (Харків)", 1600, "Па", "ДБН В.1.2-2, дод. Е"],
            ["γfm снег (T=50)", 1.00, "", "табл. 8.1"], ["γfe снег (η=0.02)", 0.49, "", "табл. 8.3"],
            ["W0 (Харків)", 430, "Па", "дод. Е"], ["γfm ветер (T=50)", 1.00, "", "табл. 9.1"],
            ["γfe ветер (η=0.02)", 0.21, "", "табл. 9.3"], ["Ch (тип III, z≤5 м)", 0.40, "", "табл. 9.01"],
            ["Ce навес тип I, 12°", "+0.62/−1.04/−0.88/−0.08", "", "схема 11 дод. И"],
            ["Cf (трение)", 0.04, "", "прим. 2 к схеме 11"],
            ["Кровля (металлочерепица)", 50, "Па по скату", "γf = 1.05 (0.95)"],
            ["Сталь, γf", 1.05, "", "табл. 5.1"], ["ψ кратковр. 1/2", "1.0 / 0.9", "", "п. 4.18"],
            ["γn (СС1)", 1.0, "", "ДБН В.1.2-14"], ["Ry", 230, "МПа", "С235/S235JR"]]
    sheet("Нагрузки", ["Параметр", "Значение", "Ед.", "Источник"], rows, [30, 26, 12, 30])
    wb.save(OUT / "smeta_optimized.xlsx")

    # ---------------- MARKDOWN ----------------
    L = []
    w = L.append
    w("# Навес из профтрубы — анализ и оптимизация каркаса\n")
    w(f"Итог по смете: **было {fmt(est_b['total'])} грн → стало {fmt(est['total'])} грн** "
      f"(экономия **{fmt(est_b['total'] - est['total'])} грн**, {100*(1-est['total']/est_b['total']):.1f} %).  ")
    w(f"Металл (трубы по метражу): {fmt(est_b['pipes'])} → {fmt(est['pipes'])} грн; масса {q_b['mass']:.0f} → {q['mass']:.0f} кг; "
      f"деталей {q_b['n_cuts']} → {q['n_cuts']}; швов {q_b['weld_len']:.0f} → {q['weld_len']:.0f} м; "
      f"окраска {q_b['area']:.0f} → {q['area']:.0f} м².\n")
    w(f"Принятая схема: **{best['label']}**.\n")
    w("## Исходные данные и допущения\n")
    for n_ in notes:
        w(f"- {n_}")
    w("- Не менялись: уклон 12°, 4 колонны в тех же точках, кровля и шаг обрешётки 350 мм, размеры в плане, "
      "коридор проезда 4.00 × 2.90 м (подкос у левой колонны срезает верхний угол коридора: у края низ на "
      f"+{sc.knee_z + 0.40 * (M.Z_TIE - sc.knee_z) / sc.knee_dx:.2f} м — выше, чем в черновике (+2.64 м))." if sc.knee_t else
      "- Не менялись: уклон 12°, 4 колонны, кровля и шаг обрешётки, план, коридор проезда 4.00 × 2.90 м (свободен полностью).")
    w("- Морозное пучение не учитывается (по заданию).\n")
    w("## Проверка исходного варианта (черновик 2026-ПР-480) по тем же нагрузкам\n")
    w("| Элемент | Сечение | Коэф. исп. | Определяющая проверка |\n|---|---|---:|---|")
    for r in mt_b:
        w(f"| {r['name']} | □{r['sec']} | {r['u']:.2f} | {r['gov']} |")
    w("\nОсновные резервы черновика: колонны, стропила, затяжки и пояса боковых ферм загружены на 30–50 %; "
      "снег принят с γfm = 1.14 (T = 100 лет) вместо 1.00 (T = 50 лет); колонны заведены в бетон на всю глубину лунки.\n")
    w("## Оптимизированный каркас\n")
    w("| Элемент | Сечение | N сж, кН | N раст, кН | M, кН·м | Коэф. исп. | Определяющая проверка |\n|---|---|---:|---:|---:|---:|---|")
    for r in mt:
        w(f"| {r['name']} | □{r['sec']} | {r['Nc']:.1f} | {r['Nt']:.1f} | {r['M']:.2f} | {r['u']:.2f} | {r['gov']} |")
    amin = min(v for v in an.alpha.values() if v)
    w(f"\nМинимальный коэффициент запаса общей устойчивости α_cr = {amin:.1f}.\n")
    w("### Опирание (лунки)\n")
    w("| Колонна | V сж, кН | V выдёрг., кН | H, кН | M, кН·м | Исп. по грунту (Бромс) |\n|---|---:|---:|---:|---:|---:|")
    for (x, y), r in an.found_results.items():
        w(f"| x={x:.2f}, y={y:.2f} | {r['Vcomp']/1e3:.1f} | {r['Vuplift']/1e3:.1f} | {r.get('H',0)/1e3:.2f} | "
          f"{r.get('M',0)/1e3:.2f} | {r['u']:.2f} |")
    w("\nДавление под пятой лунки 0.3×0.3 м от наибольшего сжатия ≈ "
      f"{max(r['Vcomp'] for r in an.found_results.values())/1e3/0.09:.0f} кПа (без учёта трения по боковой поверхности) — "
      "для суглинка тугопластичного допустимо; при слабом грунте (насыпь, текучий суглинок) расширить низ лунки.\n")
    w("### Ведомость деталей\n")
    w("| Деталь | Сечение | Длина, м | Кол-во | Масса, кг |\n|---|---|---:|---:|---:|")
    for p in pieces:
        w(f"| {p['name']} | □{p['sec']} | {p['L']:.3f} | {p['n']} | {p['mass']:.1f} |")
    w("\n### Металл по профилям\n")
    w("| Профиль | Длина, м | Масса, кг | Цена, грн/м | Сумма, грн |\n|---|---:|---:|---:|---:|")
    for name, d in sorted(est["by_sec"].items(), key=lambda kv: -kv[1]["cost"]):
        w(f"| □{name} | {d['L']:.2f} | {d['kg']:.1f} | {d['sec'].price:g} | {fmt(d['cost'])} |")
    w(f"| **Итого трубы** | | | | **{fmt(est['pipes'])}** |\n")
    w("## Смета: было → стало\n")
    w("| Статья | Было, грн | Стало, грн | Δ, грн |\n|---|---:|---:|---:|")
    for n_, v, k in est["lines"]:
        b = base_lines.get(n_, 0.0)
        w(f"| {n_} | {fmt(b)} | {fmt(v)} | {fmt(v-b)} |")
    w(f"| **Итого** | **{fmt(est_b['total'])}** | **{fmt(est['total'])}** | **{fmt(est['total']-est_b['total'])}** |\n")
    w("Статьи кровли, водостока, доставки и работ по ним не менялись. Окраска пересчитана по площади поверхности, "
      "расходники — по длине швов, изготовление и монтаж — по условной трудоёмкости (детали, стыки, швы, масса), "
      "откалиброванной так, что для черновика получаются суммы из сметы.\n")
    w("## Сравнение схем\n")
    w("| # | Схема | Итог, грн | Трубы, грн | Масса, кг |\n|---:|---|---:|---:|---:|")
    for i, r in enumerate(res_all[:15], 1):
        w(f"| {i} | {r['label']} | {fmt(r['total'])} | {fmt(r['pipes'])} | {r['mass']:.0f} |")
    w("")
    if best_norm:
        w("## Если «спеццены» окажутся недоступны\n")
        w("В прайсе есть позиции заметно дешевле рынка (< 50 грн/кг): 80×40×2,2 (150 грн/м), 80×40×2,5 (182), "
          "60×60×2,5 (184), 80×80×3 хлыст 5,9 м (280), 80×80×2,5 (285), 100×100×2,5 (348), 100×100×2,7 (381). "
          "Их наличие нужно подтвердить. Без них лучший вариант:\n")
        w(f"- {best_norm['label']}: итог **{fmt(best_norm['total'])} грн**, трубы {fmt(best_norm['pipes'])} грн, "
          f"{best_norm['mass']:.0f} кг;")
        w("- сечения: " + "; ".join(f"{M.GROUP_INFO[g][0]} □{s}" for g, s in best_norm["groups"].items()) + ".\n")
    w("## Чувствительность принятого решения\n")
    w("| Условие | Макс. коэф. исп. | Элемент | Проверка |\n|---|---:|---|---|")
    for tag, (u, g, gov) in sens.items():
        w(f"| {tag} | {u:.2f} | {M.GROUP_INFO[g][0]} | {gov} |")
    w("\nКоэффициенты > 1.0 означают, что при таком допущении решение нужно усиливать.\n")
    w("## Что проверить до закупки\n")
    w("- Тип местности по ветру (III — приміська забудова). Если вокруг открытое поле — пересчитать с Ch = 0.70.")
    w("- Наличие позиций по спеццене и фактическую толщину стенки (не менее номинала).")
    w("- Грунт в первой лунке: при насыпи или супеси — углубить / расширить лунки.")
    w("- Узлы: полный провар по контуру, катет не менее толщины стенки; заглушки на торцах труб.")
    (OUT / "report.md").write_text("\n".join(L), encoding="utf-8")
    print("готово:", est_b["total"], "→", est["total"])


if __name__ == "__main__":
    main()
